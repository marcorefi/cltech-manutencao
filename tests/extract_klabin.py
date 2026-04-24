import openpyxl, json, os

src = r'C:\Users\Marcos Lima\Meu Drive (cltechinstalacoeseautomacao@gmail.com)\CLTECH Fire\Operacional\Contratos Mensais-trimestrais\Planilhas de contrato\Planilhas de contrato 2024\KLABIN SDAI.xlsx'
wb = openpyxl.load_workbook(src, data_only=True)

all_points = []
for name in wb.sheetnames:
    ws = wb[name]
    central = 'Central 01 Manutencao' if 'CENTRAL 01' in name.upper() else 'Central 02 Portaria'
    rows = list(ws.iter_rows(values_only=True))
    for i, row in enumerate(rows):
        if i < 3:
            continue
        point_name = str(row[0]).strip() if row[0] else ''
        if not point_name or point_name.lower() == 'none':
            continue
        date_m = row[5]
        date_p = row[6]
        def fdate(v):
            if not v:
                return ''
            try:
                return v.strftime('%Y-%m-%d')
            except:
                return str(v).strip()
        all_points.append({
            'POINT NAME': point_name,
            'DEVICE TYPE': str(row[1] or '').strip(),
            'POINT TYPE': str(row[2] or '').strip(),
            'DESCRICAO NA CENTRAL': str(row[3] or '').strip(),
            'CENTRAL': central,
            'PERIODO': str(row[4] or '').strip(),
            'VISITA': 'PENDENTE',
            'DATA MANUTENCAO': fdate(date_m),
            'DATA PROX MANUTENCAO': fdate(date_p),
            'TECNICO': str(row[7] or '').strip(),
            'OBSERVACAO': str(row[8] or '').strip()
        })

print(f'Total pontos: {len(all_points)}')
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'apps-script', 'klabin-pontos.json')
with open(out, 'w', encoding='utf-8') as f:
    json.dump(all_points, f, ensure_ascii=False, indent=2)
print(f'Saved to {out}')
centrais = {}
for p in all_points:
    centrais[p['CENTRAL']] = centrais.get(p['CENTRAL'], 0) + 1
print('Por central:', centrais)
print('First 2:', json.dumps(all_points[:2], ensure_ascii=False))
